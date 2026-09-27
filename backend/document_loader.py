from pathlib import Path

import fitz


def load_pdf(path):
    document = fitz.open(path)

    pages = []

    for page_number, page in enumerate(document, start=1):
        pages.append(
            {"text": page.get_text(), "source": path.name, "page": page_number}
        )

    document.close()

    return pages


def load_docx(path):
    from docx import Document

    document = Document(path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    return [{"text": "\n".join(paragraphs), "source": path.name}]


def load_xlsx(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)

    documents = []

    for sheet in workbook.worksheets:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]

            if values:
                rows.append(" | ".join(values))

        if rows:
            documents.append(
                {"text": "\n".join(rows), "source": path.name, "sheet": sheet.title}
            )

    return documents


def load_pptx(path):
    from pptx import Presentation

    presentation = Presentation(path)

    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)

        if texts:
            slides.append(
                {"text": "\n".join(texts), "source": path.name, "slide": slide_number}
            )

    return slides


def load_text(path):
    text = path.read_text(encoding="utf-8")

    return [{"text": text, "source": path.name}]


def load_csv(path):
    import csv

    rows = []

    with open(path, "r", encoding="utf-8", newline="") as file:
        reader = csv.reader(file)

        for row in reader:
            if row:
                rows.append(" | ".join(row))

    return [{"text": "\n".join(rows), "source": path.name}]


def load_json(path):
    import json

    with open(path, "r", encoding="utf-8") as file:
        data = json.load(file)

    text = json.dumps(data, indent=2, ensure_ascii=False)

    return [{"text": text, "source": path.name}]


def load_document(file_path):
    path = Path(file_path)

    extension = path.suffix.lower()

    if extension == ".pdf":
        return load_pdf(path)

    elif extension == ".docx":
        return load_docx(path)

    elif extension == ".xlsx":
        return load_xlsx(path)

        # elif extension == ".pptx":
        # return load_pptx(path)

    elif extension in [".txt", ".md"]:
        return load_text(path)

    elif extension == ".csv":
        return load_csv(path)

    elif extension == ".json":
        return load_json(path)

    else:
        raise ValueError(f"Unsupported file type: {extension}")
